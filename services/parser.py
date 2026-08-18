"""统一文档解析适配层。

解析器只负责把输入文件转换成结构化文本和来源锚点，不调用模型、不执行插件。
旧调用方仍可使用 ``parse() -> str``；新调用方应使用 ``parse_document()``。
"""
from __future__ import annotations

from abc import ABC, abstractmethod
from dataclasses import dataclass
from html.parser import HTMLParser
from io import BytesIO
from pathlib import Path
from typing import BinaryIO, Any
import logging
import os
import re
import tempfile
import time
import zipfile
import xml.etree.ElementTree as ET

from pydantic import BaseModel, Field

from core.config_manager import config_manager

logger = logging.getLogger(__name__)


class ParserLimitError(ValueError):
    """文档解析超出确定性资源限制（页数/成员数/解压大小/超时）。"""


@dataclass(frozen=True)
class ParserLimits:
    """轻量解析器统一资源限制（方向报告 P1-2）。

    页数/成员数/解压大小/超时是确定性上限，防止超大文档或恶意压缩包
    （zip bomb）拖垮宿主进程；触发时安全截断/拒绝而非 crash。
    """

    page_limit: int = 200
    epub_max_members: int = 2000
    epub_member_max_bytes: int = 10 * 1024 * 1024        # 10MB / 成员
    epub_max_decompressed_bytes: int = 50 * 1024 * 1024  # 总计 50MB
    timeout_seconds: int = 120

    def to_dict(self) -> dict[str, int]:
        return {
            "page_limit": self.page_limit,
            "epub_max_members": self.epub_max_members,
            "epub_member_max_bytes": self.epub_member_max_bytes,
            "epub_max_decompressed_bytes": self.epub_max_decompressed_bytes,
            "timeout_seconds": self.timeout_seconds,
        }


def load_parser_limits() -> ParserLimits:
    """从 config.yaml 的 parser.* 键读取限制（默认值兜底，下限 1）。"""
    return ParserLimits(
        page_limit=max(1, config_manager.get_int("parser.page_limit", 200)),
        epub_max_members=max(1, config_manager.get_int("parser.epub_max_members", 2000)),
        epub_member_max_bytes=max(
            1, config_manager.get_int("parser.epub_member_max_bytes", 10 * 1024 * 1024)
        ),
        epub_max_decompressed_bytes=max(
            1, config_manager.get_int("parser.epub_max_decompressed_bytes", 50 * 1024 * 1024)
        ),
        timeout_seconds=max(1, config_manager.get_int("parser.timeout_seconds", 120)),
    )


class ParsedPassage(BaseModel):
    sequence: int = Field(default=0, ge=0)
    heading: str = ""
    text: str
    char_start: int | None = None
    char_end: int | None = None
    page_start: int | None = None
    page_end: int | None = None


class ParsedDocument(BaseModel):
    title: str = ""
    media_type: str = "text/plain"
    parser_id: str
    parser_version: str = "1.0"
    passages: list[ParsedPassage] = Field(default_factory=list)
    # ── Docling 可选生产接入新增（缺省为空，向后兼容） ──
    tables: list[dict[str, Any]] = Field(default_factory=list)
    pages: list[dict[str, Any]] = Field(default_factory=list)
    provenance: dict[str, Any] = Field(default_factory=dict)
    truncated: bool = False

    @property
    def text(self) -> str:
        return "\n\n".join(p.text for p in self.passages if p.text.strip())


class DocumentParser(ABC):
    """文档解析器抽象基类。"""

    parser_id = "unknown"
    parser_version = "1.0"
    media_type = "text/plain"

    @abstractmethod
    def parse_document(self, file: BinaryIO, filename: str) -> ParsedDocument:
        pass

    def parse(self, file: BinaryIO, filename: str) -> str:
        """兼容旧 API，仅返回纯文本。"""
        return self.parse_document(file, filename).text


class DoclingParser(DocumentParser):
    """可选 Docling 适配器；未安装时由工厂回退具体轻量解析器。"""
    parser_id = "docling"
    parser_version = "optional"

    def __init__(self, guard: Any = None):
        # 延迟导入：未启用 docling 时不产生任何额外开销
        from services.docling_guard import DoclingGuard, get_default_guard

        # 默认共享进程级 guard（信号量跨请求共享，实现全局并发上限）
        self.guard: DoclingGuard = guard or get_default_guard()

    def parse_document(self, file: BinaryIO, filename: str) -> ParsedDocument:
        try:
            from docling.document_converter import DocumentConverter  # type: ignore
        except ImportError as exc:
            raise ValueError("Docling 未安装") from exc
        from services.docling_guard import compute_source_hash

        # 生产约束：离线模式缺模型工件直接失败（带预下载命令提示）
        self.guard.check_offline()
        suffix = Path(filename).suffix or ".bin"
        data = file.read()
        with tempfile.NamedTemporaryFile(suffix=suffix, delete=False) as temporary:
            temporary.write(data)
            temporary_path = Path(temporary.name)
        start = time.monotonic()
        try:
            # CPU 线程限制（尽力而为：docling 版本参数不一，不支持时回退默认）
            try:
                converter = DocumentConverter(num_threads=self.guard.limits.cpu_threads)
            except TypeError:
                converter = DocumentConverter()
            result = converter.convert(str(temporary_path))
            text = result.document.export_to_markdown()
            document = ParsedDocument(
                title=Path(filename).stem,
                media_type=self.media_type,
                parser_id=self.parser_id,
                parser_version=self.parser_version,
                passages=_split_text_passages(text),
            )
            # 页数截断 + 结构化 provenance（超时/并发在 async_parse_document 中施加）
            return self.guard.apply_limits(
                document,
                self,
                source_hash=compute_source_hash(data),
                elapsed_ms=int((time.monotonic() - start) * 1000),
            )
        finally:
            temporary_path.unlink(missing_ok=True)

    async def async_parse_document(self, file: BinaryIO, filename: str) -> ParsedDocument:
        """生产入口：叠加并发信号量 + 超时 + 截断 + provenance。

        同步 parse_document 已包含离线/截断/provenance；异步版本额外提供
        asyncio.Semaphore 并发上限与解析超时保护，供可异步调用的上传入口使用。
        """
        return await self.guard.run_parse(self, file, filename)


class TxtParser(DocumentParser):
    parser_id = "builtin.txt"
    media_type = "text/plain"

    def parse_document(self, file: BinaryIO, filename: str) -> ParsedDocument:
        try:
            content = file.read()
            try:
                text = content.decode("utf-8-sig")
            except UnicodeDecodeError:
                logger.warning("文件 %s UTF-8 解码失败，尝试 GBK", filename)
                text = content.decode("gbk", errors="ignore")
        except Exception as exc:
            raise ValueError(f"无法解析文件 {filename}: {exc}") from exc
        passages = _split_text_passages(text)
        return ParsedDocument(
            title=Path(filename).stem,
            media_type="text/markdown" if filename.lower().endswith((".md", ".markdown")) else self.media_type,
            parser_id=self.parser_id,
            parser_version=self.parser_version,
            passages=passages,
        )


class PdfParser(DocumentParser):
    parser_id = "pypdf"
    media_type = "application/pdf"

    def parse_document(self, file: BinaryIO, filename: str) -> ParsedDocument:
        try:
            from pypdf import PdfReader
        except ImportError as exc:
            raise ValueError("PDF 解析需要安装 pypdf") from exc
        try:
            limits = load_parser_limits()
            reader = PdfReader(file)
            passages: list[ParsedPassage] = []
            offset = 0
            truncated = False
            for page_no, page in enumerate(reader.pages, start=1):
                # 方向报告 P1-2：页数上限，超出安全截断不再继续解压后续页
                if page_no > limits.page_limit:
                    truncated = True
                    logger.warning("PDF 超过页数上限（%d），在 %d 页截断", limits.page_limit, page_no - 1)
                    break
                text = (page.extract_text() or "").strip()
                if not text:
                    continue
                start = offset
                offset += len(text) + 2
                passages.append(ParsedPassage(
                    sequence=len(passages), text=text,
                    char_start=start, char_end=start + len(text),
                    page_start=page_no, page_end=page_no,
                ))
            return ParsedDocument(
                title=Path(filename).stem, media_type=self.media_type,
                parser_id=self.parser_id, parser_version=self.parser_version,
                passages=passages, truncated=truncated,
            )
        except Exception as exc:
            raise ValueError(f"无法解析 PDF {filename}: {exc}") from exc


class DocxParser(DocumentParser):
    parser_id = "python-docx"
    media_type = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

    def parse_document(self, file: BinaryIO, filename: str) -> ParsedDocument:
        try:
            from docx import Document
        except ImportError as exc:
            raise ValueError("DOCX 解析需要安装 python-docx") from exc
        try:
            document = Document(file)
            passages: list[ParsedPassage] = []
            offset = 0
            current_heading = ""
            for paragraph in document.paragraphs:
                text = paragraph.text.strip()
                if not text:
                    continue
                style = str(getattr(paragraph.style, "name", "") or "")
                if style.lower().startswith("heading"):
                    current_heading = text
                    continue
                start = offset
                offset += len(text) + 2
                passages.append(ParsedPassage(
                    sequence=len(passages), heading=current_heading, text=text,
                    char_start=start, char_end=start + len(text),
                ))
            return ParsedDocument(
                title=Path(filename).stem, media_type=self.media_type,
                parser_id=self.parser_id, parser_version=self.parser_version,
                passages=passages,
            )
        except Exception as exc:
            raise ValueError(f"无法解析 DOCX {filename}: {exc}") from exc


class _HTMLTextExtractor(HTMLParser):
    _heading_tags = {"h1", "h2", "h3", "h4", "h5", "h6"}

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.parts: list[str] = []
        self.current_heading = ""
        self.in_heading = False

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        self.in_heading = tag.lower() in self._heading_tags

    def handle_endtag(self, tag: str) -> None:
        if tag.lower() in self._heading_tags:
            self.in_heading = False
            self.parts.append("\n")

    def handle_data(self, data: str) -> None:
        text = re.sub(r"\s+", " ", data).strip()
        if not text:
            return
        if self.in_heading:
            self.parts.append(f"\n## {text}\n")
        else:
            self.parts.append(text + " ")


class EpubParser(DocumentParser):
    parser_id = "builtin.epub-zip"
    media_type = "application/epub+zip"

    def parse_document(self, file: BinaryIO, filename: str) -> ParsedDocument:
        try:
            with zipfile.ZipFile(file, "r") as archive:
                limits = load_parser_limits()
                names = [name for name in archive.namelist() if name.lower().endswith((".xhtml", ".html", ".htm"))]

                # 方向报告 P1-2：EPUB 归档成员数上限（防解压数量爆炸）
                total_members = len(archive.namelist())
                if total_members > limits.epub_max_members:
                    raise ParserLimitError(
                        f"EPUB 归档成员数 {total_members} 超过上限 {limits.epub_max_members}，拒绝解析"
                    )

                passages: list[ParsedPassage] = []
                offset = 0
                total_decompressed = 0

                for name in sorted(names):
                    # 方向报告 P1-2：单成员解压后字节上限（防单文件膨胀）
                    info = archive.getinfo(name)
                    if info.file_size > limits.epub_member_max_bytes:
                        logger.warning(
                            "EPUB 成员 %s 解压后 %d 字节超过上限 %d，跳过该成员",
                            name, info.file_size, limits.epub_member_max_bytes,
                        )
                        continue
                    # 全归档解压累计字节上限（防 zip bomb）
                    total_decompressed += info.file_size
                    if total_decompressed > limits.epub_max_decompressed_bytes:
                        logger.warning(
                            "EPUB 累计解压 %d 字节超过上限 %d，截断剩余成员",
                            total_decompressed, limits.epub_max_decompressed_bytes,
                        )
                        break

                    # 只读取归档成员，不解压到文件系统；避免路径穿越和临时文件污染。
                    raw = archive.read(name)
                    extractor = _HTMLTextExtractor()
                    extractor.feed(raw.decode("utf-8", errors="ignore"))
                    text = re.sub(r"\n{3,}", "\n\n", "".join(extractor.parts)).strip()
                    if not text:
                        continue
                    start = offset
                    offset += len(text) + 2
                    passages.append(ParsedPassage(
                        sequence=len(passages), heading=Path(name).stem, text=text,
                        char_start=start, char_end=start + len(text),
                    ))
                return ParsedDocument(
                    title=Path(filename).stem, media_type=self.media_type,
                    parser_id=self.parser_id, parser_version=self.parser_version,
                    passages=passages,
                )
        except ParserLimitError:
            raise
        except zipfile.BadZipFile as exc:
            raise ValueError(f"EPUB 不是有效 ZIP 文档: {filename}") from exc
        except Exception as exc:
            raise ValueError(f"无法解析 EPUB {filename}: {exc}") from exc


def _split_text_passages(text: str) -> list[ParsedPassage]:
    normalized = text.replace("\r\n", "\n").replace("\r", "\n")
    blocks = [block.strip() for block in re.split(r"\n\s*\n", normalized) if block.strip()]
    if not blocks and normalized.strip():
        blocks = [normalized.strip()]
    passages: list[ParsedPassage] = []
    cursor = 0
    for sequence, block in enumerate(blocks):
        start = normalized.find(block, cursor)
        if start < 0:
            start = cursor
        end = start + len(block)
        heading = ""
        first_line = block.splitlines()[0].strip() if block.splitlines() else ""
        if re.match(r"^(#{1,6}\s+|第[^\n]{1,30}[章节卷]|Chapter\s+)" , first_line, re.I):
            heading = first_line.lstrip("# ")
        passages.append(ParsedPassage(
            sequence=sequence, heading=heading, text=block,
            char_start=start, char_end=end,
        ))
        cursor = end
    return passages


class ParserFactory:
    """按扩展名选择真实解析器；未知格式不静默伪装成结构化文档。"""

    @staticmethod
    def get_parser(filename: str) -> DocumentParser:
        ext = Path(filename).suffix.lower()
        if ext in ("", ".txt", ".md", ".markdown"):
            return TxtParser()
        # Docling 作为可选增强：默认关闭（docling.enabled=false 回退轻量解析器）。
        # 启用条件：config.yaml 设置 docling.enabled=true，或环境变量
        # AIILIE_DOCLING_ENABLED=true（config_manager 自动映射），且 docling 可导入。
        if ext in (".pdf", ".docx", ".epub"):
            use_docling = config_manager.get_bool("docling.enabled", False)
            # 兼容旧环境变量开关 AIILIE_PARSER_DOCLING=1
            if os.environ.get("AIILIE_PARSER_DOCLING", "0") == "1":
                use_docling = True
            if use_docling:
                try:
                    import docling  # type: ignore  # noqa: F401
                    return DoclingParser()
                except ImportError:
                    logger.warning("docling.enabled=true 但 Docling 未安装，回退轻量解析器")
            return {".pdf": PdfParser, ".docx": DocxParser, ".epub": EpubParser}[ext]()
        if ext == ".doc":
            raise ValueError("旧式 .doc 暂不支持，请转换为 .docx 或 PDF 后导入")
        raise ValueError(f"不支持的文档格式: {ext or '(无扩展名)'}")
